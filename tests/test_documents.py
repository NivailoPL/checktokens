import io
import zipfile

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

from checktokens.core import count_file
from checktokens.extract import extract_docx, extract_pdf, extract_rtf


def pdf_bytes():
    stream = io.BytesIO()
    canvas = Canvas(stream)
    canvas.drawString(30, 700, "hello world")
    canvas.save()
    return stream.getvalue()


def test_rtf_text():
    assert extract_rtf(rb"{\rtf1\ansi hello world}").text == "hello world"
    assert extract_rtf(rb"{\rtf1\ansi\ansicpg1250 za\'bf\'f3\'b3\'e6}").text == "zażółć"


def test_pdf_text_and_empty_page():
    assert extract_pdf(pdf_bytes()).text.strip() == "hello world"
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf_bytes())))
    writer.add_blank_page(width=600, height=800)
    stream = io.BytesIO()
    writer.write(stream)
    extracted = extract_pdf(stream.getvalue())
    assert "hello world" in extracted.text
    assert extracted.warnings == ["1 of 2 pages have no extracted text; OCR may be needed."]


def test_docx_order_nested_tables_and_headers():
    doc = Document()
    doc.add_paragraph("before")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "nested"
    doc.add_paragraph("after")
    doc.sections[0].header.paragraphs[0].text = "header"
    doc.add_section()
    stream = io.BytesIO()
    doc.save(stream)
    text = extract_docx(stream.getvalue()).text
    assert text.index("before") < text.index("left") < text.index("right")
    assert text.index("nested") < text.index("after")
    assert text.count("header") == 1


def test_docx_unsupported_content_warning():
    doc = Document()
    doc.add_paragraph("body")
    stream = io.BytesIO()
    doc.save(stream)
    out = io.BytesIO()
    with zipfile.ZipFile(stream) as source, zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "word/document.xml":
                data = data.replace(
                    b"</w:body>", b"<w:ins><w:p><w:r><w:t>change</w:t></w:r></w:p></w:ins></w:body>"
                )
            target.writestr(item, data)
    result = extract_docx(out.getvalue())
    assert result.warnings
    assert "change" not in result.text


def test_password_and_no_text_pdf(tmp_path):
    path = tmp_path / "document.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.write(path)
    assert "No text" in count_file(str(path)).error
    writer.encrypt("secret")
    writer.write(path)
    assert "password" in count_file(str(path)).error


def test_corrupt_docx(tmp_path):
    path = tmp_path / "document.docx"
    path.write_bytes(b"not a docx")
    assert count_file(str(path)).error
